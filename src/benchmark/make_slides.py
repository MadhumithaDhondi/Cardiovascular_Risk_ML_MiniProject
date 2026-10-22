"""Build the combined review deck (Study A + Study B) from results/*. Run AFTER train.py + visualize.py for BOTH datasets.

    python src/benchmark/make_slides.py --names "Name1, Name2"
"""
import argparse
import json
from pathlib import Path

import pandas as pd
from PIL import Image as PILImage
from pptx import Presentation
from pptx.util import Inches, Pt

import datasets

REPO = datasets.REPO


def load(R):
    R = Path(R)
    return {"sum": json.load(open(R / "data_summary.json")), "m": pd.read_csv(R / "metrics.csv"),
            "sub": pd.read_csv(R / "subgroups.csv"), "cal": json.load(open(R / "calibration.json")), "fig": R / "figures"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cardio", default=str(REPO / "results" / "cardio"))
    ap.add_argument("--framingham", default=str(REPO / "results" / "framingham"))
    ap.add_argument("--out", default=str(REPO / "slides" / "deck.pptx"))
    ap.add_argument("--names", default="<Team members>")
    a = ap.parse_args()
    A, B = load(a.cardio), load(a.framingham)
    sa, sb = A["sum"], B["sum"]
    bestA = A["m"].sort_values("test_auroc").iloc[-1]
    bestB = B["m"].sort_values("test_auroc").iloc[-1]

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)

    def set_title(sl, text, size=30):
        t = sl.shapes.title
        t.text = text
        t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.0)
        for p in t.text_frame.paragraphs:
            p.font.size = Pt(size)
            p.font.bold = True

    def bullets(title, items):
        sl = prs.slides.add_slide(prs.slide_layouts[1])
        set_title(sl, title)
        body = sl.placeholders[1]
        body.left, body.top, body.width, body.height = Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.5)
        tf = body.text_frame
        tf.word_wrap = True
        for i, t in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = t
            p.font.size = Pt(22)

    def picture(title, path, caption=None):
        sl = prs.slides.add_slide(prs.slide_layouts[5])
        set_title(sl, title, 26)
        w, h = PILImage.open(path).size
        width = min(12.3, 5.0 * w / h)
        sl.shapes.add_picture(str(path), Inches((13.33 - width) / 2), Inches(1.6), width=Inches(width))
        if caption:
            tb = sl.shapes.add_textbox(Inches(0.6), Inches(6.6), Inches(12.1), Inches(0.6))
            tb.text_frame.text = caption
            tb.text_frame.paragraphs[0].font.size = Pt(16)

    sl = prs.slides.add_slide(prs.slide_layouts[0])
    sl.shapes.title.text = "Predicting Cardiovascular Risk using Electronic Health Records"
    sl.placeholders[1].text = f"UE24CS352A Machine Learning - Mini-Project\n{a.names}"
    for ph, top, hgt in [(sl.shapes.title, 2.2, 1.8), (sl.placeholders[1], 4.3, 1.5)]:
        ph.left, ph.top, ph.width, ph.height = Inches(1.2), Inches(top), Inches(10.9), Inches(hgt)

    bullets("Problem and plan", [
        "CVD is the leading cause of death worldwide - early identification of at-risk patients matters",
        "Question 1: which models predict cardiovascular risk best from routine clinical data?",
        "Question 2: do the conclusions hold on two very different datasets?",
        "Following Nguyen (Stanford CS229); the original STRIDE EHR data is private, so we use two public datasets",
        "Study A: Kaggle (prevalent CVD).  Study B: Framingham (10-year incident CHD)"])
    bullets("Two studies, one pipeline", [
        f"Study A - Kaggle: {sa['n_raw']:,} records -> {sa['n_clean']:,} clean; {sa['n_features']} features; {sa['positive_rate']*100:.0f}% positive (balanced)",
        f"Study B - Framingham: {sb['n_raw']:,} records; {sb['n_features']} features; {sb['positive_rate']*100:.0f}% positive (imbalanced); "
        f"{sb['n_missing_rows']:,} rows with missing values",
        "Same five models, same split, same CV, same metrics on both",
        "Framingham exploration and first baselines done earlier; the benchmark reuses the exact same split"])
    picture("Data quality - Study A (Kaggle)", A["fig"] / "fig_data_quality.png", "Implausible blood pressure / height / weight rows removed; no missing values")
    picture("Data quality - Study B (Framingham)", B["fig"] / "fig_data_quality.png", "Missing values imputed with the training-fold median (inside the pipeline, no leakage)")
    picture("Exploratory analysis - Study A", A["fig"] / "fig_eda.png")
    picture("Exploratory analysis - Study B", B["fig"] / "fig_eda.png")
    bullets("Approach and protocol", [
        "Logistic regression: Ridge, Lasso, Elastic Net (standardised features)",
        "Random Forest: bagged, decorrelated trees",
        "XGBoost: sequential regularised trees, each correcting the previous errors",
        f"Stratified 80/20 split (seed 42); {sa['cv_folds']}-fold CV grid search on AUROC, training data only",
        "Imputation and scaling inside each fold; class weighting for the 15% positive class (Study B)",
        "Metrics: AUROC, avg precision, sensitivity/specificity, confusion matrix, calibration, subgroup AUROC"])
    bullets("Implementation", [
        "src/benchmark/datasets.py - loading, cleaning and settings per dataset",
        "src/benchmark/train.py - tuning and evaluation; visualize.py - all figures",
        "src/benchmark/make_report.py / make_slides.py - build this write-up and deck from results",
        "src/*.py - Framingham exploration and first baselines; results/ - outputs for both studies",
        "Python: pandas, scikit-learn, XGBoost, matplotlib"])

    sl = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(sl, "Results: test AUROC")
    ms = A["m"].merge(B["m"], on="model", suffixes=("_a", "_b"))
    tbl = sl.shapes.add_table(len(ms) + 1, 5, Inches(0.8), Inches(1.7), Inches(11.7), Inches(0.6) * (len(ms) + 1)).table
    for j, hd in enumerate(["Model", "A: CV AUROC", "A: Test AUROC", "B: CV AUROC", "B: Test AUROC"]):
        tbl.cell(0, j).text = hd
    for i, r in enumerate(ms.itertuples(), start=1):
        for j, v in enumerate([r.model, f"{r.cv_auroc_a:.3f}", f"{r.test_auroc_a:.3f}", f"{r.cv_auroc_b:.3f}", f"{r.test_auroc_b:.3f}"]):
            tbl.cell(i, j).text = v
    picture("Model comparison - Study A", A["fig"] / "fig_models.png")
    picture("Model comparison - Study B", B["fig"] / "fig_models.png",
            f"{sb['n_test']} test patients: differences of a few hundredths in AUROC are within noise")
    picture("Confusion matrices - Study A", A["fig"] / "fig_confusion.png", "Threshold 0.5; percentages are row-normalised")
    picture("Confusion matrices - Study B", B["fig"] / "fig_confusion.png", "Class weighting trades specificity for sensitivity")
    picture("Importance, subgroups, calibration - Study A", A["fig"] / "fig_diagnostics.png")
    picture("Importance, subgroups, calibration - Study B", B["fig"] / "fig_diagnostics.png",
            f"Class weighting inflates probabilities (mean {B['cal']['mean_predicted']:.2f} vs observed {B['cal']['observed_rate']:.2f})")
    spreadA = A["m"].test_auroc.max() - A["m"].test_auroc.min()
    spreadB = B["m"].test_auroc.max() - B["m"].test_auroc.min()
    gapAB = abs(bestA.test_auroc - bestB.test_auroc)
    lower = ("AUROC is lower for 10-year incident risk on a small cohort (Study B) than for prevalent disease (Study A)"
             if bestB.test_auroc < bestA.test_auroc else "AUROC is similar or higher in Study B than in Study A")
    bullets("Conclusions", [
        f"Study A best: {bestA.model}, test AUROC {bestA.test_auroc:.3f}.  Study B best: {bestB.model}, test AUROC {bestB.test_auroc:.3f}",
        f"Spread across the five models: {spreadA:.3f} (A) and {spreadB:.3f} (B), vs a {gapAB:.3f} gap between the datasets - "
        + ("the data matters more than the model" if gapAB > max(spreadA, spreadB) else "model choice matters as much as the data"),
        lower,
        "Performance is not uniform across age groups - audit before clinical use",
        "Class weighting keeps the ranking (AUROC) but distorts probabilities - recalibrate before use"])
    bullets("Limitations and future work", [
        "Study A: prevalent (not incident) disease, noisy self-reported fields, no ethnicity column",
        "Study B: small single cohort, one random split",
        "Next: repeated CV and confidence intervals, probability calibration",
        "Next: survival-style modelling of incident risk; fairness-aware training"])
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    prs.save(a.out)
    print("wrote", a.out)


if __name__ == "__main__":
    main()
